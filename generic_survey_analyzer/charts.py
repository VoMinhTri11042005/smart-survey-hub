"""
Reusable, Excel-safe chart + table helpers built on openpyxl.

IMPORTANT — lessons learned the hard way (see PITFALLS.md):
  1. Category `Reference` ranges must exclude the header row (top+1 .. bottom),
     never (top .. top) or (top .. bottom) including the header, or bars/pies
     shift by one and show the header text as a phantom category/series.
  2. For a clustered "from_rows=True" chart, the data Reference must start at
     top+1 (labels+values only); the header row is used ONLY as `categories`.
  3. `DataLabelList` must be assigned to BOTH `chart.dataLabels` and
     `series.dLbls` (the same object) or LibreOffice falls back to printing
     "SeriesName; value" next to every point.
  4. NEVER set `axis.title = ''`. Assigning an empty string still creates a
     title element and some renderers print the literal text "None". Only
     assign a title when the string is non-empty.
  5. Do NOT run a "recalculate formulas" pass (LibreOffice `store()`) on a
     workbook that has zero formulas — it re-serialises every chart in
     LibreOffice's own dialect and can trigger Excel's "we found a problem
     with some content" repair prompt. Skip recalculation entirely when the
     workbook has no formulas (this report never uses formulas — every
     number is computed in Python and written as a static value).
"""
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter

from style import PALETTE, BLUE, FONT_HDR, FONT_BODY, FILL_HDR, BORDER
from openpyxl.styles import Alignment


def style_header_row(ws, row, col_start, col_end):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = FONT_HDR
        cell.fill = FILL_HDR
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER


def write_table(ws, start_row, start_col, headers, rows, col_widths=None):
    """Writes a header + data block, returns the row *after* the last data row."""
    for j, h in enumerate(headers):
        ws.cell(row=start_row, column=start_col + j, value=h)
    style_header_row(ws, start_row, start_col, start_col + len(headers) - 1)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = ws.cell(row=start_row + 1 + i, column=start_col + j, value=v)
            cell.font = FONT_BODY
            cell.border = BORDER
            if j > 0:
                cell.alignment = Alignment(horizontal='center')
    if col_widths:
        for j, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(start_col + j)].width = w
    return start_row + 1 + len(rows)


def color_points(series_obj, n, palette=PALETTE):
    pts = []
    for i in range(n):
        dp = DataPoint(idx=i)
        dp.graphicalProperties = GraphicalProperties(solidFill=palette[i % len(palette)])
        pts.append(dp)
    series_obj.data_points = pts


def add_bar_chart(ws, anchor, title, cat_ref, min_col, max_col, min_row, max_row,
                   bar_dir='col', color=BLUE, y_title='', x_title='', vary_colors=False,
                   height=8.5, width=17):
    """cat_ref = (cat_col_start, cat_col_end, header_row, last_data_row)."""
    chart = BarChart()
    chart.type = bar_dir
    chart.style = 10
    chart.title = title
    if y_title:
        chart.y_axis.title = y_title
    if x_title:
        chart.x_axis.title = x_title
    chart.height = height
    chart.width = width
    chart.legend = None
    chart.gapWidth = 60

    data = Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row, max_row=max_row)
    cat_col_start, cat_col_end, cat_top, cat_bottom = cat_ref
    cats = Reference(ws, min_col=cat_col_start, max_col=cat_col_end,
                      min_row=cat_top + 1, max_row=cat_bottom)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    s = chart.series[0]
    n_pts = max_row - min_row
    if vary_colors:
        color_points(s, n_pts)
    else:
        s.graphicalProperties.solidFill = color

    dl = DataLabelList()
    dl.showVal = True
    dl.showSerName = False
    dl.showCatName = False
    dl.showLegendKey = False
    chart.dataLabels = dl
    s.dLbls = dl
    ws.add_chart(chart, anchor)
    return chart


def add_pie_chart(ws, anchor, title, cat_ref, min_col, max_col, min_row, max_row,
                   doughnut=False, hole_size=55, height=9, width=13, show_pct=True):
    chart = DoughnutChart() if doughnut else PieChart()
    chart.title = title
    chart.height = height
    chart.width = width
    if doughnut:
        chart.holeSize = hole_size

    data = Reference(ws, min_col=min_col, max_col=max_col, min_row=min_row, max_row=max_row)
    cat_col_start, cat_col_end, cat_top, cat_bottom = cat_ref
    cats = Reference(ws, min_col=cat_col_start, max_col=cat_col_end,
                      min_row=cat_top + 1, max_row=cat_bottom)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    s = chart.series[0]
    color_points(s, max_row - min_row)

    dl = DataLabelList()
    dl.showVal = False
    dl.showPercent = show_pct
    dl.showCatName = False
    dl.showSerName = False
    dl.showLegendKey = False
    dl.showBubbleSize = False
    chart.dataLabels = dl
    s.dLbls = dl
    chart.legend.position = 'r'
    ws.add_chart(chart, anchor)
    return chart


def add_clustered_from_rows(ws, anchor, title, table_top, table_bottom, n_cat_cols,
                             label_col, y_title='', height=10, width=22):
    """
    Build a clustered column chart where each ROW of the table is a series
    (e.g. one per top category) and the columns to its right are the groups
    (e.g. one per demographic segment).

    `label_col` MUST equal the `start_col` you passed to `write_table()` for
    this exact table — i.e. the column holding the row labels (and, in the
    header row, holding the table's first header cell). Hardcoding column 1
    here is a classic bug when tables are written starting at column 2 (to
    leave a left margin column) — the header cell gets pulled in as a bogus
    extra category. Always pass the real label_col explicitly.

    table_top = header row of the table, table_bottom = last data row.
    """
    chart = BarChart()
    chart.type = 'col'
    chart.grouping = 'clustered'
    chart.overlap = -10
    chart.style = 10
    chart.title = title
    if y_title:
        chart.y_axis.title = y_title
    chart.height = height
    chart.width = width
    cats = Reference(ws, min_col=label_col + 1, max_col=label_col + n_cat_cols,
                      min_row=table_top, max_row=table_top)
    data = Reference(ws, min_col=label_col, max_col=label_col + n_cat_cols,
                      min_row=table_top + 1, max_row=table_bottom)
    chart.add_data(data, titles_from_data=True, from_rows=True)
    chart.set_categories(cats)
    for i, s in enumerate(chart.series):
        s.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
    ws.add_chart(chart, anchor)
    return chart
