"""Design constants shared by every chart/sheet in the report."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY, BLUE, ACCENT, GREEN, PURPLE, ORANGE, GOLD, TEAL, GREY2 = \
    '1F3864', '2E5FA3', 'C00000', '548235', '7030A0', 'ED7D31', 'FFC000', '2E9C9C', '7F7F7F'
PALETTE = [BLUE, ACCENT, GREEN, ORANGE, PURPLE, GOLD, TEAL, NAVY, GREY2, '4472C4']

FONT_TITLE = Font(name='Arial', size=16, bold=True, color=NAVY)
FONT_SUB = Font(name='Arial', size=10, italic=True, color='595959')
FONT_SEC = Font(name='Arial', size=12, bold=True, color=NAVY)
FONT_HDR = Font(name='Arial', size=11, bold=True, color='FFFFFF')
FONT_BODY = Font(name='Arial', size=10)
FONT_INSIGHT = Font(name='Arial', size=10, italic=True, color='375623')
FILL_HDR = PatternFill('solid', fgColor=NAVY)
FILL_KPI = PatternFill('solid', fgColor='E9EFF7')
FILL_INSIGHT = PatternFill('solid', fgColor='EDF7ED')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
