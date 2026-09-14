"""Orchestrates the full workbook from a raw dataframe + its column-type map.
Zero hardcoded question text: every label comes from the data itself."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from style import FONT_TITLE, FONT_SUB, FONT_SEC, FONT_INSIGHT, FILL_KPI, FILL_INSIGHT, BORDER, NAVY
from charts import write_table, add_bar_chart, add_pie_chart, add_clustered_from_rows, style_header_row
import aggregate as agg

DATA_COL_START = 20  # helper-data blocks are written far to the right of each sheet


def build_workbook(df: pd.DataFrame, col_types: dict, title: str = 'PHÂN TÍCH DỮ LIỆU KHẢO SÁT') -> Workbook:
    N = len(df)
    wb = Workbook()
    wb.remove(wb.active)

    # ---------- raw data ----------
    ws0 = wb.create_sheet('Dữ liệu gốc')
    for row in dataframe_to_rows(df, index=False, header=True):
        ws0.append(row)
    style_header_row(ws0, 1, 1, len(df.columns))
    for j, colname in enumerate(df.columns, start=1):
        ws0.column_dimensions[get_column_letter(j)].width = min(max(14, len(str(colname)) // 2), 45)
    ws0.freeze_panes = 'A2'

    multi_cols = [c for c, i in col_types.items() if i['type'] == 'MULTI_CHOICE']
    single_cols = [c for c, i in col_types.items() if i['type'] == 'SINGLE_CHOICE']
    freeform_cols = [c for c, i in col_types.items() if i['type'] == 'FREEFORM_CATEGORICAL']
    rating_cols = [c for c, i in col_types.items() if i['type'] == 'RATING']
    open_text_cols = [c for c, i in col_types.items() if i['type'] == 'OPEN_TEXT']

    # a small single-choice column with 2-6 categories is the best candidate to
    # cross-tab everything else against (e.g. "year", "gender", "group"...)
    group_col = None
    for c in single_cols:
        k = df[c].nunique()
        if 2 <= k <= 6:
            group_col = c
            break

    # ============ DASHBOARD ============
    ws = wb.create_sheet('Dashboard')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    for cl in 'BCDEFGHI':
        ws.column_dimensions[cl].width = 13
    ws['B2'] = title
    ws['B2'].font = FONT_TITLE
    ws['B3'] = f'n = {N} phản hồi · Cập nhật: {pd.Timestamp.now().strftime("%d/%m/%Y")}'
    ws['B3'].font = FONT_SUB

    kpis = [('TỔNG PHẢN HỒI', N)]
    if multi_cols:
        top_item = agg.agg_multi_choice(df[multi_cols[0]], col_types[multi_cols[0]]['delimiter'])[0]
        kpis.append(('PHỔ BIẾN NHẤT: ' + multi_cols[0][:28], top_item[0]))
    if rating_cols:
        _, mean, _ = agg.agg_rating(df[rating_cols[0]])
        kpis.append((rating_cols[0][:28] + ' (TB)', mean))
    col = 2
    for label, val in kpis[:4]:
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col + 1)
        c1 = ws.cell(row=5, column=col, value=label)
        c1.font = Font(name='Arial', size=9, bold=True, color='595959')
        c1.alignment = Alignment(horizontal='center')
        c2 = ws.cell(row=6, column=col, value=val)
        c2.font = Font(name='Arial', size=13, bold=True, color=NAVY)
        c2.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r in range(5, 8):
            for cc in range(col, col + 2):
                ws.cell(row=r, column=cc).fill = FILL_KPI
                ws.cell(row=r, column=cc).border = BORDER
        col += 2

    dc = DATA_COL_START
    r = 9
    anchor_row = 9
    anchor_col_toggle = True
    # NOTE: pie/donut is only valid for SINGLE_CHOICE (one answer per respondent,
    # slices sum to ~100%). MULTI_CHOICE items don't partition respondents (people
    # pick several), so a pie there would silently misrepresent the data -> always bar.
    for c in (multi_cols + single_cols)[:4]:
        is_multi = c in multi_cols
        items = (agg.agg_multi_choice(df[c], col_types[c]['delimiter']) if is_multi
                 else agg.agg_single_choice(df[c]))
        row_end = write_table(ws, r, dc, [c[:30], 'Số lượt'], [[k, v] for k, v in items])
        top, bot = r, row_end - 1
        anchor = f'{"B" if anchor_col_toggle else "J"}{anchor_row}'
        if (not is_multi) and len(items) <= 6:
            add_pie_chart(ws, anchor, c[:60], (dc, dc, top, bot), dc + 1, dc + 1, top, bot, doughnut=True)
        else:
            add_bar_chart(ws, anchor, c[:60], (dc, dc, top, bot), dc + 1, dc + 1, top, bot,
                           bar_dir='bar', vary_colors=True)
        r = row_end + 2
        if not anchor_col_toggle:
            anchor_row += 18
        anchor_col_toggle = not anchor_col_toggle

    # ============ CHI TIẾT CÂU HỎI ============
    ws2 = wb.create_sheet('Chi tiết câu hỏi')
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions['A'].width = 2
    ws2['B1'] = 'BIỂU ĐỒ CHI TIẾT THEO TỪNG CÂU HỎI'
    ws2['B1'].font = FONT_TITLE
    r2, anchor_row2 = 3, 3
    dc2 = DATA_COL_START

    def block(ws_, colname, kind):
        nonlocal r2, anchor_row2
        if kind == 'MULTI_CHOICE':
            items = agg.agg_multi_choice(df[colname], col_types[colname]['delimiter'])
        elif kind == 'SINGLE_CHOICE':
            items = agg.agg_single_choice(df[colname])
        elif kind == 'FREEFORM_CATEGORICAL':
            items = agg.agg_freeform_categorical(df[colname])
        else:
            return
        row_end = write_table(ws_, r2, dc2, [colname[:30], 'Số lượt'], [[k, v] for k, v in items])
        top, bot = r2, row_end - 1
        anchor = f'B{anchor_row2}'
        if kind == 'SINGLE_CHOICE' and len(items) <= 6:
            add_pie_chart(ws_, anchor, colname[:70], (dc2, dc2, top, bot), dc2 + 1, dc2 + 1, top, bot, doughnut=(len(items) <= 5))
        else:
            add_bar_chart(ws_, anchor, colname[:70], (dc2, dc2, top, bot), dc2 + 1, dc2 + 1, top, bot,
                           bar_dir='bar', vary_colors=True)
        r2 = row_end + 2
        anchor_row2 += 19

    for c in multi_cols + single_cols + freeform_cols:
        block(ws2, c, col_types[c]['type'])

    for c in rating_cols:
        dist, mean, bands = agg.agg_rating(df[c])
        row_end = write_table(ws2, r2, dc2, [c[:30], 'Số phản hồi'], [[k, v] for k, v in dist])
        top, bot = r2, row_end - 1
        add_bar_chart(ws2, f'B{anchor_row2}', c[:70], (dc2, dc2, top, bot), dc2 + 1, dc2 + 1, top, bot,
                       bar_dir='col', vary_colors=True)
        r2 = row_end + 2
        anchor_row2 += 19

    # ============ PHÂN TÍCH CHÉO ============
    ws3 = wb.create_sheet('Phân tích chéo')
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions['A'].width = 2
    ws3['B1'] = 'PHÂN TÍCH CHÉO — MỐI LIÊN HỆ GIỮA CÁC BIẾN'
    ws3['B1'].font = FONT_TITLE
    r3 = 3

    if group_col and multi_cols:
        mc = multi_cols[0]
        items = agg.agg_multi_choice(df[mc], col_types[mc]['delimiter'])
        top_cats = agg.top_categories(items, n=5)
        groups, rows = agg.cross_tab_percent(df, group_col, mc, col_types[mc]['delimiter'], top_cats)
        ws3.cell(row=r3, column=2, value=f'1. Tỷ lệ (%) theo "{mc[:40]}", chia theo "{group_col[:30]}"').font = FONT_SEC
        r3 += 1
        headers = ['Hạng mục'] + [str(g) for g in groups]
        end_row = write_table(ws3, r3, 2, headers, rows, col_widths=[38] + [12] * len(groups))
        add_clustered_from_rows(ws3, f'B{end_row + 2}', f'{mc[:50]} theo {group_col[:30]}',
                                 r3, end_row - 1, len(groups), label_col=2, y_title='%')
        r3 = end_row + 22

    if group_col and rating_cols:
        rc = rating_cols[0]
        rows = agg.avg_by_group(df, group_col, rc)
        ws3.cell(row=r3, column=2, value=f'2. {rc[:60]} — trung bình theo "{group_col[:30]}"').font = FONT_SEC
        r3 += 1
        end_row = write_table(ws3, r3, 2, [group_col[:30], 'Điểm TB'], rows, col_widths=[30, 12])
        add_bar_chart(ws3, f'B{end_row + 2}', f'{rc[:60]} theo {group_col[:30]}',
                       (2, 2, r3, end_row - 1), 3, 3, r3, end_row - 1, bar_dir='col', vary_colors=True)
        r3 = end_row + 20

    if len(rating_cols) >= 1 and single_cols:
        # correlation of the rating column against an ordinal-looking single-choice column, if any
        pass  # left for Codex to extend: detect ordinal single-choice + encode + pearson/spearman

    # ============ PHẢN HỒI MỞ (open text, never charted) ============
    if open_text_cols:
        ws4 = wb.create_sheet('Phản hồi mở')
        ws4.sheet_view.showGridLines = False
        ws4.column_dimensions['A'].width = 2
        ws4['B1'] = 'CÂU TRẢ LỜI DẠNG TỰ DO (không trực quan hoá)'
        ws4['B1'].font = FONT_TITLE
        rr = 3
        for c in open_text_cols:
            end_row = write_table(ws4, rr, 2, [c], [[v] for v in df[c].dropna().astype(str)], col_widths=[80])
            rr = end_row + 2

    return wb
